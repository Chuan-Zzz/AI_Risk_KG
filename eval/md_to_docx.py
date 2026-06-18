"""将 eval_report.md 转换为格式化的 Word 文档 (.docx)"""

from __future__ import annotations

import re
import sys
import io
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

MD_PATH = Path(__file__).resolve().parent / "eval_report.md"
DOCX_PATH = Path(__file__).resolve().parent / "eval_report_v2.docx"


def set_cell_shading(cell, color: str):
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color,
        qn('w:val'): 'clear',
    })
    shading.append(shd)


def add_formatted_runs(paragraph, text: str):
    bold_pattern = re.compile(r'\*\*(.+?)\*\*')
    code_pattern = re.compile(r'`([^`]+)`')
    combined = re.compile(r'(\*\*(.+?)\*\*|`([^`]+)`)')

    pos = 0
    for m in combined.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            run.font.size = Pt(10)

        if m.group(2):
            run = paragraph.add_run(m.group(2))
            run.bold = True
            run.font.size = Pt(10)
        elif m.group(3):
            run = paragraph.add_run(m.group(3))
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)

        pos = m.end()

    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.font.size = Pt(10)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        line = line.strip()
        if not line or set(line) <= {'|', '-', ':', ' '}:
            continue
        cells = [c.strip() for c in line.split('|')]
        cells = [c for c in cells if c != '']
        rows.append(cells)
    return rows


def add_table_to_doc(doc: Document, rows: list[list[str]]):
    if not rows:
        return

    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            if j >= n_cols:
                continue
            cell = table.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            clean = cell_text.strip()
            clean = re.sub(r'\*\*(.+?)\*\*', r'\1', clean)
            clean = re.sub(r'`([^`]+)`', r'\1', clean)

            run = p.add_run(clean)
            run.font.size = Pt(9)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_cell_shading(cell, '4472C4')

    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph()


def add_code_block(doc: Document, code: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.5)

    pf = p.paragraph_format
    pf_format = p._element.get_or_add_pPr()

    shd = pf_format.makeelement(qn('w:shd'), {
        qn('w:fill'): 'F5F5F5',
        qn('w:val'): 'clear',
    })
    pf_format.append(shd)

    for line in code.split('\n'):
        if p.text == '' and not p.runs:
            run = p.add_run(line)
        else:
            run = p.add_run('\n' + line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def convert_md_to_docx(md_path: Path, docx_path: Path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    for level in range(1, 5):
        h_style = doc.styles[f'Heading {level}']
        h_style.font.name = 'Microsoft YaHei'
        h_style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^---+\s*$', line):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run('─' * 60)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            run.font.size = Pt(8)
            i += 1
            continue

        # Headings
        heading_match = re.match(r'^(#{1,4})\s+(.+)$', line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'`([^`]+)`', r'\1', text)

            h = doc.add_heading(text, level=level)
            for run in h.runs:
                run.font.name = 'Microsoft YaHei'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
            i += 1
            continue

        # Code block
        if line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            i += 1  # skip closing ```
            add_code_block(doc, '\n'.join(code_lines))
            continue

        # Table
        if '|' in line and i + 1 < len(lines) and re.match(r'^[\s|:-]+$', lines[i + 1].strip()):
            table_lines = []
            while i < len(lines) and '|' in lines[i]:
                table_lines.append(lines[i].rstrip('\n'))
                i += 1
            rows = parse_table(table_lines)
            add_table_to_doc(doc, rows)
            continue

        # Blockquote
        if line.startswith('>'):
            text = line.lstrip('> ').strip()
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'`([^`]+)`', r'\1', text)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            run.italic = True
            i += 1
            continue

        # Numbered list
        num_match = re.match(r'^(\d+)\.\s+(.+)$', line)
        if num_match:
            text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            add_formatted_runs(p, text)
            i += 1
            continue

        # Bullet list
        bullet_match = re.match(r'^[-*]\s+(.+)$', line)
        if bullet_match:
            text = bullet_match.group(1)
            p = doc.add_paragraph(style='List Bullet')
            add_formatted_runs(p, text)
            i += 1
            continue

        # Normal paragraph
        text = line.strip()
        if text:
            p = doc.add_paragraph()
            add_formatted_runs(p, text)

        i += 1

    # Set page margins
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    doc.save(str(docx_path))
    print(f"DOCX saved to: {docx_path}")


if __name__ == '__main__':
    convert_md_to_docx(MD_PATH, DOCX_PATH)
